import csv
import json
import unicodedata
from pathlib import Path
from .common import digest, dumps


def read_records(path, text_column="text", id_column="id", context_column="context", language_column="language", source_column="source", limit=None):
    path = Path(path)
    if limit is not None and limit < 1:
        raise ValueError("--limit must be positive")
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
    elif path.suffix.lower() in (".jsonl", ".ndjson"):
        rows = [json.loads(line) for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip()]
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            iterator = wb.active.iter_rows()
            headers = [c.value for c in next(iterator)]
            if len(headers) != len(set(headers)):
                raise ValueError("Duplicate XLSX column headers")
            rows = []
            for cells in iterator:
                if any(c.data_type == "f" for c in cells):
                    raise ValueError("Input workbook contains formulas; provide literal text")
                if any(c.value is not None for c in cells):
                    rows.append(dict(zip(headers, [c.value for c in cells])))
        finally:
            wb.close()
    else:
        raise ValueError("Input must be CSV, JSONL or XLSX")
    result, ids = [], set()
    for n, row in enumerate(rows[:limit], 1):
        if not isinstance(row, dict) or not isinstance(row.get(text_column), str) or not row[text_column].strip():
            raise ValueError(f"Row {n}: nonempty string column '{text_column}' required")
        text = row[text_column]
        context, language, source = (row.get(c) or "" for c in (context_column, language_column, source_column))
        if not all(isinstance(v, str) for v in (context, language, source)):
            raise ValueError(f"Row {n}: context/language/source must be strings")
        content_hash = digest({"text": text, "context": context, "language": language})
        ident = str(row.get(id_column)) if row.get(id_column) not in (None, "") else f"row-{n}-{content_hash[:12]}"
        if ident in ids:
            raise ValueError(f"Duplicate input id at row {n}: {ident}")
        ids.add(ident)
        try:
            dumps(row)
        except (ValueError, TypeError) as e:
            raise ValueError(f"Row {n}: metadata is not JSON serializable (use ISO dates)") from e
        result.append({"id": ident, "text": text, "text_normalized": unicodedata.normalize("NFC", text),
                       "context": context, "language": language, "source": source, "content_hash": content_hash,
                       "input_row": n, "metadata": row})
    if not result:
        raise ValueError("Input contains no records")
    return result
