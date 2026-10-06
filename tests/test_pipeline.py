import csv
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from offense_judge.aggregation import aggregate, cache_key, quote_hint
from offense_judge.common import atomic_jsonl, dumps
from offense_judge.config import load_config
from offense_judge.demo import DemoProvider, demo_config, fixture_decision, run_demo
from offense_judge.export import MAIN_COLUMNS, import_reviews
from offense_judge.inputs import read_records
from offense_judge.pipeline import run_pipeline
from offense_judge.provider import APIError, BudgetExceeded, Gate, build_payload, extract_content
from offense_judge.schema import InvalidResponse, decision_warnings, parse_content, validate
from offense_judge.storage import Store


def record(case="criticism", text="متن آزمایش", ident="x"):
    return {"id": ident, "text": text, "context": "", "language": "fa", "source": "test", "metadata": {"case": case},
            "content_hash": "hash-" + ident, "input_row": 1, "text_normalized": text}


def votes(labels, r=None):
    r = r or record()
    result = []
    for i, label in enumerate(labels, 1):
        d = fixture_decision(r, i)
        d.update(label=label, p_offensive_raw=None if label is None else .96 if label else .04,
                 abuse_types=["insult"] if label else [])
        result.append({"slot": i, "model": f"m{i}", "status": "ok", "decision": d, "warnings": []})
    return result


class Decisions(unittest.TestCase):
    def setUp(self):
        self.cfg = demo_config()
        self.r = record()
        self.d = fixture_decision(self.r, 1)

    def test_prefix_and_fenced_json(self):
        text = "دلیل کوتاه\nFINAL_JSON:\n```json\n" + dumps(self.d) + "\n```"
        d, reason = parse_content(text)
        self.assertEqual(d, self.d)
        self.assertEqual(reason, "دلیل کوتاه")

    def test_think_json_is_not_a_vote(self):
        d, reason = parse_content("<think>{\"label\":1}</think>\nدلیل\nFINAL_JSON:" + dumps(self.d))
        self.assertEqual(d["label"], 0)

    def test_ambiguous_multiple_objects_fail(self):
        with self.assertRaises(InvalidResponse):
            parse_content(dumps(self.d) + "\n" + dumps(self.d))

    def test_duplicate_json_key_rejected(self):
        with self.assertRaises(InvalidResponse):
            parse_content(dumps(self.d).replace('"label": 0', '"label": 1, "label": 0'))

    def test_boolean_and_string_labels_fail(self):
        for label in (True, False, "0", 2):
            with self.subTest(label=label), self.assertRaises(InvalidResponse):
                validate({**self.d, "label": label})

    def test_bad_probability_and_unknown_keys_fail(self):
        for p in (float("nan"), float("inf"), -1, 2, "0.5", True):
            with self.subTest(p=p), self.assertRaises(InvalidResponse):
                validate({**self.d, "p_offensive_raw": p})
        with self.assertRaises(InvalidResponse):
            validate({**self.d, "extra": "x"})

    def test_nan_json_rejected(self):
        with self.assertRaises(InvalidResponse):
            parse_content(dumps(self.d).replace('0.04', 'NaN'))

    def test_refusal_or_truncation_not_valid(self):
        for reason in ("length", "content_filter", None):
            with self.subTest(reason=reason), self.assertRaises(InvalidResponse):
                extract_content({"choices": [{"finish_reason": reason, "message": {"content": dumps(self.d)}}]}, "chat")
        with self.assertRaises(InvalidResponse):
            extract_content({"choices": [{"finish_reason": "stop", "message": {"content": dumps(self.d), "refusal": "refused"}}]}, "chat")

    def test_responses_typed_output(self):
        content, reasoning = extract_content({"status": "completed", "output": [{"type": "message", "content": [{"type": "output_text", "text": dumps(self.d)}]}]}, "responses")
        self.assertEqual(parse_content(content)[0], self.d)
        with self.assertRaises(InvalidResponse):
            extract_content({"status": "incomplete", "output_text": dumps(self.d)}, "responses")
        with self.assertRaises(InvalidResponse):
            extract_content({"status": "completed", "output": [{"content": [{"type": "refusal", "refusal": "no"}]}]}, "responses")

    def test_three_votes_are_candidate_and_default_silver(self):
        r = aggregate(self.r, votes([1, 1, 1, 0]), self.cfg.policy, "run")
        self.assertEqual(r["candidate_label"], 1)
        self.assertEqual(r["final_label"], 1)
        self.assertIn("minority_vote", r["review_flags"])

    def test_tie_abstention_missing_never_zero(self):
        for vv in (votes([1, 0, 1, 0]), votes([0, 0, 0, None]), votes([0, 0, 0])):
            r = aggregate(self.r, vv, self.cfg.policy, "run")
            self.assertIsNone(r["final_label"])
            self.assertIsNone(r["candidate_label"])
            self.assertTrue(r["needs_review"])

    def test_failed_fourth_is_not_zero(self):
        vv = votes([0, 0, 0, 0])
        vv[3].update(status="error", decision=None)
        result = aggregate(self.r, vv, self.cfg.policy, "run")
        self.assertEqual(result["status"], "error")
        self.assertIsNone(result["final_label"])

    def test_quotation_blocked_even_unanimous(self):
        r = record(text='او گفت «سلام».')
        result = aggregate(r, votes([0, 0, 0, 0], r), self.cfg.policy, "run")
        self.assertEqual(result["candidate_label"], 0)
        self.assertIsNone(result["final_label"])
        self.assertIn("review_quotation", result["review_flags"])
        self.assertFalse(quote_hint("It's fine, don't worry."))

    def test_educational_flag_blocks(self):
        vv = votes([0, 0, 0, 0])
        vv[0]["decision"]["discourse_tags"] = ["educational"]
        self.assertIsNone(aggregate(self.r, vv, self.cfg.policy, "run")["final_label"])

    def test_context_low_confidence_blocks(self):
        vv = votes([0, 0, 0, 0])
        vv[1]["decision"]["p_offensive_raw"] = .45
        result = aggregate(self.r, vv, self.cfg.policy, "run")
        self.assertIsNone(result["final_label"])
        self.assertIn("low_vote_confidence_raw", result["review_flags"])

    def test_positive_subtype_disagreement_blocks(self):
        vv = votes([1, 1, 1, 0])
        vv[1]["decision"]["abuse_types"] = ["hate_speech"]
        r = aggregate(self.r, vv, self.cfg.policy, "run")
        self.assertEqual(r["candidate_label"], 1)
        self.assertIsNone(r["final_label"])
        self.assertIn("subtype_disagreement", r["review_flags"])

    def test_evidence_must_be_in_target_for_positive(self):
        d = fixture_decision(record("insult"), 1)
        d["evidence"] = [{"text": "فقط زمینه", "kind": "abuse"}]
        r = {**self.r, "context": "فقط زمینه"}
        warnings, locations = decision_warnings(d, r)
        self.assertIn("missing_target_text_abuse_evidence", warnings)
        self.assertEqual(locations[0]["matches"][0]["field"], "context")

    def test_cache_depends_on_prompt_context_model_not_key(self):
        a = cache_key(self.cfg, self.cfg.models[0], self.r)
        self.assertNotEqual(a, cache_key(self.cfg, self.cfg.models[0], {**self.r, "context": "new"}))
        self.cfg.key = "SECRET"
        self.assertEqual(a, cache_key(self.cfg, self.cfg.models[0], self.r))
        self.cfg.prompt += "edited"
        self.assertNotEqual(a, cache_key(self.cfg, self.cfg.models[0], self.r))

    def test_route_parameters(self):
        m = self.cfg.models[0]
        p = build_payload(self.cfg, m, self.r)
        self.assertNotIn("response_format", p)
        from dataclasses import replace
        p = build_payload(self.cfg, replace(m, api="responses", output_mode="json_schema", reasoning_effort="low"), self.r)
        self.assertEqual(p["text"]["format"]["type"], "json_schema")
        self.assertEqual(p["reasoning"], {"effort": "low"})
        self.assertNotIn("messages", p)

    def test_budget_is_threadsafe_and_counts_calls(self):
        gate = Gate(max_calls=1)
        gate.acquire()
        with self.assertRaises(BudgetExceeded):
            gate.acquire()

    def test_persian_normalization_in_evidence_matching(self):
        d = fixture_decision(record("insult", text="این متن شامل بی شعور است"), 1)
        d["evidence"] = [{"text": "بي‌شعور", "kind": "abuse"}]
        r = record("insult", text="این متن شامل بی شعور است")
        warnings, locations = decision_warnings(d, r)
        self.assertNotIn("evidence_not_found", warnings)

    def test_finish_reason_end_turn(self):
        raw = {"choices": [{"finish_reason": "end_turn", "message": {"content": dumps(self.d)}}]}
        content, _ = extract_content(raw, "chat")
        self.assertEqual(parse_content(content)[0], self.d)

    def test_moderation_tiebreak_resolution(self):
        tie_votes = votes([1, 1, 0, 0])
        res_no_mod = aggregate(self.r, tie_votes, self.cfg.policy, "run")
        self.assertIsNone(res_no_mod["candidate_label"])
        self.assertEqual(res_no_mod["status"], "tie")

        mod_pos = {"status": "ok", "decision": {"flagged": True, "label": 1, "flagged_categories": ["harassment"]}}
        res_pos = aggregate(self.r, tie_votes, self.cfg.policy, "run", moderation=mod_pos)
        self.assertEqual(res_pos["candidate_label"], 1)
        self.assertEqual(res_pos["final_label"], 1)
        self.assertEqual(res_pos["label_origin"], "moderation_tiebreak")
        self.assertIn("tie_2_2_moderation_resolved", res_pos["review_flags"])

        mod_neg = {"status": "ok", "decision": {"flagged": False, "label": 0, "flagged_categories": []}}
        res_neg = aggregate(self.r, tie_votes, self.cfg.policy, "run", moderation=mod_neg)
        self.assertEqual(res_neg["candidate_label"], 0)
        self.assertEqual(res_neg["final_label"], 0)
        self.assertEqual(res_neg["label_origin"], "moderation_tiebreak")
        self.assertIn("tie_2_2_moderation_resolved", res_neg["review_flags"])

    def test_two_judge_aggregation_and_1_1_moderation_resolution(self):
        policy_2 = {**self.cfg.policy, "num_judges": 2}
        v_unanimous = votes([1, 1])
        res_unanimous = aggregate(self.r, v_unanimous, policy_2, "run", expected_judges=2)
        self.assertEqual(res_unanimous["candidate_label"], 1)
        self.assertEqual(res_unanimous["final_label"], 1)

        v_tie = votes([1, 0])
        res_tie = aggregate(self.r, v_tie, policy_2, "run", expected_judges=2)
        self.assertIsNone(res_tie["candidate_label"])
        self.assertEqual(res_tie["status"], "tie")

        mod_pos = {"status": "ok", "decision": {"flagged": True, "label": 1, "flagged_categories": ["harassment"]}}
        res_pos = aggregate(self.r, v_tie, policy_2, "run", moderation=mod_pos, expected_judges=2)
        self.assertEqual(res_pos["candidate_label"], 1)
        self.assertEqual(res_pos["final_label"], 1)
        self.assertEqual(res_pos["label_origin"], "moderation_tiebreak")
        self.assertIn("tie_1_1_moderation_resolved", res_pos["review_flags"])

        mod_neg = {"status": "ok", "decision": {"flagged": False, "label": 0, "flagged_categories": []}}
        res_neg = aggregate(self.r, v_tie, policy_2, "run", moderation=mod_neg, expected_judges=2)
        self.assertEqual(res_neg["candidate_label"], 0)
        self.assertEqual(res_neg["final_label"], 0)
        self.assertEqual(res_neg["label_origin"], "moderation_tiebreak")
        self.assertIn("tie_1_1_moderation_resolved", res_neg["review_flags"])

    def test_three_judge_aggregation(self):
        policy_3 = {**self.cfg.policy, "num_judges": 3}
        res_3_0 = aggregate(self.r, votes([1, 1, 1]), policy_3, "run", expected_judges=3)
        self.assertEqual(res_3_0["candidate_label"], 1)
        self.assertNotIn("minority_vote", res_3_0["review_flags"])

        res_2_1 = aggregate(self.r, votes([1, 1, 0]), policy_3, "run", expected_judges=3)
        self.assertEqual(res_2_1["candidate_label"], 1)
        self.assertIn("minority_vote", res_2_1["review_flags"])

        res_1_2 = aggregate(self.r, votes([1, 0, 0]), policy_3, "run", expected_judges=3)
        self.assertEqual(res_1_2["candidate_label"], 0)
        self.assertIn("minority_vote", res_1_2["review_flags"])

    def test_atomic_jsonl_creates_nested_parent_dirs(self):
        from offense_judge.common import now
        nested = Path(tempfile.gettempdir()) / f"test_atomic_{now().replace(':', '')}" / "sub" / "data.jsonl"
        try:
            atomic_jsonl(nested, [{"test": 123}])
            self.assertTrue(nested.exists())
            self.assertEqual(json.loads(nested.read_text().strip()), {"test": 123})
        finally:
            if nested.exists():
                nested.unlink()
            if nested.parent.exists():
                nested.parent.rmdir()
            if nested.parent.parent.exists():
                nested.parent.parent.rmdir()


class EndToEnd(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_resume_duplicate_and_manifest_no_key(self):
        cfg = demo_config()
        cfg.key = "VERY_SECRET_KEY"
        records = [record(ident="a"), record(ident="b")]
        provider = DemoProvider()
        summary, _ = run_pipeline(records, cfg, self.root / "run", provider, progress=lambda _: None)
        self.assertEqual(provider.gate.calls, 4)
        run_pipeline(records, cfg, self.root / "run", provider, progress=lambda _: None)
        self.assertEqual(provider.gate.calls, 4)
        self.assertNotIn(cfg.key, (self.root / "run/manifest.json").read_text())
        with self.assertRaises(ValueError):
            run_pipeline([{**records[0], "context": "changed"}], cfg, self.root / "run", provider, progress=lambda _: None)

    def test_budget_partial_export_then_resume(self):
        cfg = demo_config()
        provider = DemoProvider()
        provider.gate = Gate(max_calls=2)
        with patch("offense_judge.pipeline.time.sleep"):
            summary, stopped = run_pipeline([record()], cfg, self.root / "run", provider, progress=lambda _: None)
        self.assertIsNotNone(stopped)
        self.assertEqual(summary["final_label_counts"], {"None": 1})
        next_provider = DemoProvider()
        summary, stopped = run_pipeline([record()], cfg, self.root / "run", next_provider, progress=lambda _: None)
        self.assertEqual(next_provider.gate.calls, 2)
        self.assertIsNone(stopped)
        self.assertEqual(summary["final_label_counts"], {"0": 1})

    def test_invalid_response_retry_saved(self):
        class BadOnce(DemoProvider):
            def call(self, model, r, repair=False):
                raw = super().call(model, r, repair)
                if not repair:
                    raw["choices"][0]["message"]["content"] = "{invalid}"
                return raw
        with patch("offense_judge.pipeline.time.sleep"):
            summary, _ = run_pipeline([record()], demo_config(), self.root / "run", BadOnce(), progress=lambda _: None)
        self.assertEqual(summary["api_attempts"], 8)
        attempts = [json.loads(line) for line in (self.root / "run/raw_responses.jsonl").read_text().splitlines()]
        self.assertEqual(sum(a["status"] == "error" for a in attempts), 4)

    def test_auth_failure_stops_and_is_preserved(self):
        class Unauthorized(DemoProvider):
            def call(self, model, r, repair=False):
                raise APIError("HTTP 401", status=401, raw="Unauthorized")
        summary, stopped = run_pipeline([record()], demo_config(), self.root / "run", Unauthorized(), progress=lambda _: None)
        self.assertEqual(stopped, "HTTP 401")
        self.assertIsNone(json.loads((self.root / "run/results.jsonl").read_text())["final_label"])
        self.assertGreater(summary["api_attempts"], 0)

    def test_failed_votes_require_explicit_retry(self):
        cfg = demo_config()
        r = record("error")
        provider = DemoProvider()
        run_pipeline([r], cfg, self.root / "run", provider, progress=lambda _: None)
        initial = provider.gate.calls
        run_pipeline([r], cfg, self.root / "run", provider, progress=lambda _: None)
        self.assertEqual(provider.gate.calls, initial)
        run_pipeline([r], cfg, self.root / "run", provider, retry_failed=True, progress=lambda _: None)
        self.assertEqual(provider.gate.calls, initial + 1)

    def test_long_text_stays_in_json_and_review_can_import(self):
        r = record(text="متن " * 9000)
        output = self.root / "run"
        run_pipeline([r], demo_config(), output, DemoProvider(), progress=lambda _: None)
        self.assertEqual(json.loads((output / "results.jsonl").read_text())["text"], r["text"])
        wb = load_workbook(output / "results.xlsx")
        ws = wb["Results"]
        heads = [c.value for c in ws[1]]
        ws.cell(2, heads.index("human_label") + 1, 0)
        ws.cell(2, heads.index("reviewer") + 1, "tester")
        copy = output / "reviewed.xlsx"
        wb.save(copy)
        wb.close()
        self.assertEqual(import_reviews(output, copy), 1)

    def test_demo_colors_formula_safety_and_review_zero(self):
        output = self.root / "demo"
        summary = run_demo(output)
        self.assertEqual(summary["real_api_calls"], 0)
        self.assertEqual(summary["api_attempts"], 44)
        wb = load_workbook(output / "results.xlsx")
        ws = wb["Results"]
        headers = [c.value for c in ws[1]]
        by_id = {ws.cell(n, 1).value: n for n in range(2, ws.max_row + 1)}
        self.assertEqual(ws.cell(by_id["fa-05"], 1).fill.fgColor.rgb[-6:], "FFF2A8")
        self.assertEqual(ws.cell(by_id["fa-03"], 1).fill.fgColor.rgb[-6:], "FFE3BA")
        self.assertEqual(ws.cell(by_id["en-03"], 1).fill.fgColor.rgb[-6:], "FFD4D4")
        self.assertEqual(ws.cell(by_id["fa-09"], 2).data_type, "s")
        n = by_id["fa-03"]
        ws.cell(n, headers.index("human_label") + 1, 0)
        ws.cell(n, headers.index("reviewer") + 1, "tester")
        copy = output / "reviewed.xlsx"
        wb.save(copy)
        wb.close()
        self.assertEqual(import_reviews(output, copy), 1)
        results = [json.loads(line) for line in (output / "results.jsonl").read_text().splitlines()]
        r = next(r for r in results if r["id"] == "fa-03")
        self.assertEqual(r["final_label"], 0)
        self.assertEqual(r["status"], "human_reviewed")
        self.assertIn("review_quotation", r["review_flags"])
        self.assertFalse(r["needs_review"])

    def test_foreign_review_rejected_before_any_write(self):
        output = self.root / "demo"
        run_demo(output)
        wb = load_workbook(output / "results.xlsx")
        ws = wb["Results"]
        headers = [c.value for c in ws[1]]
        ws.cell(2, headers.index("human_label") + 1, 1)
        ws.cell(2, headers.index("reviewer") + 1, "tester")
        ws.cell(3, headers.index("run_id") + 1, "another-run")
        copy = output / "reviewed.xlsx"
        wb.save(copy)
        wb.close()
        with self.assertRaises(ValueError):
            import_reviews(output, copy)
        store = Store(output)
        self.assertEqual(store.reviews(), {})
        store.close()

    def test_input_preservation_and_duplicate_ids(self):
        path = self.root / "input.csv"
        path.write_text('id,text,context\na,"دو خط\nداده",زمینه\nb,"دو خط\nداده",زمینه\n', encoding="utf-8")
        rows = read_records(path)
        self.assertEqual(rows[0]["text"], "دو خط\nداده")
        self.assertEqual(rows[0]["content_hash"], rows[1]["content_hash"])
        path.write_text("id,text\na,hello\na,world\n")
        with self.assertRaises(ValueError):
            read_records(path)

    def test_config_four_distinct_and_extra_guard(self):
        root = Path(__file__).resolve().parent.parent
        text = (root / ".env.example").read_text().replace("PROMPT_FILE=prompts/classifier.md", f"PROMPT_FILE={root / 'prompts/classifier.md'}").replace("REVIEW_POLICY_FILE=config/review_policy.json", f"REVIEW_POLICY_FILE={root / 'config/review_policy.json'}")
        for i in range(1, 5):
            text = text.replace(f"MODEL_{i}=\n", f"MODEL_{i}=model-{i}\n")
        env = self.root / ".env"
        env.write_text(text)
        self.assertEqual(len(load_config(env, require_key=False).models), 4)
        env.write_text(text.replace("MODEL_4=model-4", "MODEL_4=model-1"))
        with self.assertRaises(ValueError):
            load_config(env, require_key=False)
        env.write_text(text.replace("MODEL_1_EXTRA_JSON={}", 'MODEL_1_EXTRA_JSON={"model":"override"}'))
        with self.assertRaises(ValueError):
            load_config(env, require_key=False)

    def test_num_judges_env_setting(self):
        root = Path(__file__).resolve().parent.parent
        base_text = (root / ".env.example").read_text().replace("PROMPT_FILE=prompts/classifier.md", f"PROMPT_FILE={root / 'prompts/classifier.md'}").replace("REVIEW_POLICY_FILE=config/review_policy.json", f"REVIEW_POLICY_FILE={root / 'config/review_policy.json'}")
        for i in range(1, 5):
            base_text = base_text.replace(f"MODEL_{i}=\n", f"MODEL_{i}=model-{i}\n")
        env = self.root / ".env_num_judges"
        
        # Test 2 judges
        env.write_text(base_text + "\nNUM_JUDGES=2\n")
        cfg2 = load_config(env, require_key=False)
        self.assertEqual(len(cfg2.models), 2)
        self.assertEqual([m.id for m in cfg2.models], ["model-1", "model-2"])
        self.assertEqual(cfg2.policy.get("num_judges"), 2)

        # Test 3 judges
        env.write_text(base_text + "\nNUM_JUDGES=3\n")
        cfg3 = load_config(env, require_key=False)
        self.assertEqual(len(cfg3.models), 3)
        self.assertEqual([m.id for m in cfg3.models], ["model-1", "model-2", "model-3"])
        self.assertEqual(cfg3.policy.get("num_judges"), 3)

        # Test invalid judges
        env.write_text(base_text + "\nNUM_JUDGES=5\n")
        with self.assertRaises(ValueError):
            load_config(env, require_key=False)

        env.write_text(base_text + "\nNUM_JUDGES=1\n")
        with self.assertRaises(ValueError):
            load_config(env, require_key=False)

    def test_moderation_tiebreak_e2e(self):
        cfg = demo_config()
        cfg.moderation_model = "omni-moderation-latest"
        r = record(case="tie", text="چه نابغه‌ای! باز هم همه‌چیز را خراب کردی.")
        output = self.root / "tie_run"
        summary, stopped = run_pipeline([r], cfg, output, DemoProvider(), progress=lambda _: None)
        self.assertIsNone(stopped)
        results = [json.loads(line) for line in (output / "results.jsonl").read_text().splitlines()]
        self.assertEqual(len(results), 1)
        res = results[0]
        self.assertEqual(res["candidate_label"], 1)
        self.assertEqual(res["final_label"], 1)
        self.assertEqual(res["label_origin"], "moderation_tiebreak")
        self.assertIn("tie_2_2_moderation_resolved", res["review_flags"])
        self.assertIsNotNone(res.get("moderation"))
        self.assertTrue(res["moderation"]["flagged"])
        wb = load_workbook(output / "results.xlsx")
        ws = wb["Results"]
        headers = [c.value for c in ws[1]]
        self.assertIn("moderation_label", headers)
        self.assertIn("moderation_flagged", headers)
        self.assertIn("moderation_categories", headers)
        wb.close()

    def test_two_judge_moderation_tiebreak_e2e(self):
        cfg = demo_config()
        cfg.models = cfg.models[:2]
        cfg.policy["num_judges"] = 2
        cfg.moderation_model = "omni-moderation-latest"

        class TwoJudgeTieProvider(DemoProvider):
            def call(self, model, record, repair=False):
                raw = super().call(model, record, repair)
                slot = 1 if model.slot == 1 else 3
                d = fixture_decision(record, slot)
                raw["choices"][0]["message"]["content"] = dumps(d)
                return raw

        r = record(ident="tie2", case="tie", text="چه نابغه‌ای! باز هم همه‌چیز را خراب کردی.")
        output = self.root / "tie2_run"
        summary, stopped = run_pipeline([r], cfg, output, TwoJudgeTieProvider(), progress=lambda _: None)
        self.assertIsNone(stopped)
        results = [json.loads(line) for line in (output / "results.jsonl").read_text().splitlines()]
        self.assertEqual(len(results), 1)
        res = results[0]
        self.assertEqual(res["candidate_label"], 1)
        self.assertEqual(res["final_label"], 1)
        self.assertEqual(res["label_origin"], "moderation_tiebreak")
        self.assertIn("tie_1_1_moderation_resolved", res["review_flags"])


if __name__ == "__main__":
    unittest.main()
