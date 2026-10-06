import csv
import io
import os
import re
import tempfile
from collections import Counter
from pathlib import Path
from .common import atomic_jsonl, atomic_text, dumps

COLORS = {"auto_accepted": "F3F6FA", "review_required": "FFE3BA", "tie": "FFF2A8", "ambiguous": "FFF2A8", "error": "FFD4D4", "human_reviewed": "D7EEDD"}
MAIN_COLUMNS = ["id", "text", "candidate_label", "final_label", "needs_review", "status", "review_flags",
                "human_label", "reviewer", "review_note", "language", "context", "source", "positive_votes",
                "negative_votes", "blocking_review", "abuse_types_union", "discourse_tags_union", "label_origin",
                "quality_tier", "tag_annotation_origin", "duplicate_of", "content_hash", "run_id",
                "moderation_label", "moderation_flagged", "moderation_categories"]
MAIN_COLUMNS += [f"m{i}_{field}" for i in range(1, 5) for field in ("model", "label", "p_offensive_raw", "vote_confidence_raw", "reason", "status")]
VOTE_COLUMNS = ["id", "slot", "model", "status", "label", "p_offensive_raw", "vote_confidence_raw", "abuse_types", "discourse_tags",
                "decision_reason", "reasoning_before_json", "evidence", "warnings", "error", "score_type", "target_types",
                "target_basis", "tag_annotation_origin", "last_attempt_id", "cache_key", "usage"]


def flat_value(value):
    return dumps(value) if isinstance(value, (dict, list)) else value


def excel_text(value):
    value = re.sub(r"[\x00-\x08\x0b-\x0c\x0e-\x1f\ufffe\uffff]", "", value)
    return value[:32650] + " … [full value in JSONL]" if len(value) > 32700 else value


def main_row(record):
    result = {k: record.get(k) for k in MAIN_COLUMNS}
    mod = record.get("moderation") or {}
    result["moderation_label"] = mod.get("label")
    result["moderation_flagged"] = mod.get("flagged")
    result["moderation_categories"] = ",".join(mod.get("flagged_categories", [])) if mod.get("flagged_categories") else ""
    for v in record["votes"]:
        d = v.get("decision") or {}
        slot = v["slot"]
        p, label = d.get("p_offensive_raw"), d.get("label")
        confidence = None if p is None or label is None else p if label == 1 else 1 - p
        result.update({f"m{slot}_model": v["model"], f"m{slot}_label": label, f"m{slot}_p_offensive_raw": p,
                       f"m{slot}_vote_confidence_raw": confidence, f"m{slot}_reason": d.get("decision_reason", ""),
                       f"m{slot}_status": v["status"]})
    return result


def vote_rows(results):
    rows = []
    for r in results:
        for v in r["votes"]:
            d = v.get("decision") or {}
            p, label = d.get("p_offensive_raw"), d.get("label")
            rows.append({"id": r["id"], **v, **d, "vote_confidence_raw": None if p is None or label is None else p if label == 1 else 1 - p})
    return rows


def write_csv(path, rows, columns):
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        safe = {}
        for k in columns:
            value = flat_value(row.get(k))
            # Prevent spreadsheet formula execution. Original text remains exact in JSONL.
            if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
                value = "'" + value
            safe[k] = value
        writer.writerow(safe)
    atomic_text(path, "\ufeff" + stream.getvalue())


def export_run(store, results, manifest):
    directory = store.directory
    attempts_count = 0
    usage = Counter()
    for a in store.iter_attempts():
        attempts_count += 1
        u = a.get("usage") or {}
        for key in ("prompt_tokens", "completion_tokens", "input_tokens", "output_tokens", "total_tokens"):
            if type(u.get(key)) is int:
                usage[key] += u[key]
        p_details = u.get("prompt_tokens_details") or {}
        if type(p_details.get("cached_tokens")) is int:
            usage["cached_prompt_tokens"] += p_details["cached_tokens"]
        c_details = u.get("completion_tokens_details") or {}
        if type(c_details.get("reasoning_tokens")) is int:
            usage["reasoning_tokens"] += c_details["reasoning_tokens"]
    summary = {"run_id": manifest["run_id"], "records": len(results), "status_counts": dict(Counter(r["status"] for r in results)),
               "candidate_label_counts": dict(Counter(str(r["candidate_label"]) for r in results)),
               "final_label_counts": dict(Counter(str(r["final_label"]) for r in results)),
               "review_records": sum(r["needs_review"] for r in results),
               "blocking_records": sum(r["blocking_review"] for r in results),
               "moderation_tiebreak_records": sum(any(f.startswith("tie_") and f.endswith("_moderation_resolved") for f in r.get("review_flags", [])) for r in results),
               "duplicate_records": sum(bool(r["duplicate_of"]) for r in results),
               "api_attempts": attempts_count, "usage_totals_as_reported": dict(usage),
               "estimated_cost": None, "note": "Raw scores are not calibrated. Token sums include retries; duplicate records reuse requests. No pricing assumed."}
    rows, votes = [main_row(r) for r in results], vote_rows(results)
    for filename, data in (("results.jsonl", results), ("votes.jsonl", votes), ("raw_responses.jsonl", store.iter_attempts()), ("review_log.jsonl", store.review_log())):
        atomic_jsonl(directory / filename, data)
    atomic_text(directory / "summary.json", dumps(summary, indent=2))
    write_csv(directory / "results.csv", rows, MAIN_COLUMNS)
    write_csv(directory / "review_queue.csv", [row for row, r in zip(rows, results) if r["needs_review"]], MAIN_COLUMNS)
    write_csv(directory / "clean.csv", [row for row, r in zip(rows, results) if r["final_label"] is not None], MAIN_COLUMNS)
    write_workbook(directory / "results.xlsx", rows, votes, summary)
    return summary


def write_workbook(path, rows, votes, summary):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    # 20k records create >2 million cells: streaming avoids retaining them in RAM.
    wb = Workbook(write_only=True)
    main = wb.create_sheet("Results")
    main.sheet_view.rightToLeft = True
    voting = wb.create_sheet("Votes")
    info = wb.create_sheet("Summary")
    truncations = 0
    body_font = Font(name="Calibri", size=11)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_text = Alignment(vertical="top", wrap_text=True)
    align_number = Alignment(vertical="top", horizontal="center")
    fills = {key: PatternFill("solid", fgColor=color) for key, color in COLORS.items()}
    neutral_fill = PatternFill("solid", fgColor="FFFFFF")

    def put(ws, columns, data, count):
        nonlocal truncations
        ws.freeze_panes = "C2" if ws.title != "Summary" else "A2"
        ws.sheet_view.showGridLines = False
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, count + 1)}"
        header = []
        for col, name in enumerate(columns, 1):
            width = 60 if name == "text" else 45 if name in ("context", "review_flags", "decision_reason", "reasoning_before_json") or name.endswith("_reason") else 25 if name in ("id", "review_note", "source", "status") or name.endswith("_model") else 20
            if ws.title == "Summary":
                width = 34 if col == 1 else 105
            ws.column_dimensions[get_column_letter(col)].width = width
            if ws.title == "Results" and name in ("content_hash", "run_id"):
                ws.column_dimensions[get_column_letter(col)].hidden = True
            cell = WriteOnlyCell(ws, value=name)
            editable = ws.title == "Results" and name in ("human_label", "reviewer", "review_note")
            cell.fill = PatternFill("solid", fgColor="0C756E" if editable else "173B56")
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if editable:
                cell.comment = Comment("Editable review field. Import via review command; editing alone does not update final_label.", "Pipeline")
            header.append(cell)
        ws.row_dimensions[1].height = 34
        ws.append(header)
        for n, row in enumerate(data, 2):
            cells = []
            for name in columns:
                value = flat_value(row.get(name))
                cell = WriteOnlyCell(ws)
                if isinstance(value, str):
                    if len(value) > 32700:
                        truncations += 1
                    cell.value = excel_text(value)
                    cell.data_type = "s"
                    cell.alignment = align_text
                else:
                    cell.value = value
                    cell.alignment = align_number
                cell.font = body_font
                cell.fill = fills.get(row.get("status"), neutral_fill)
                if "confidence" in name or "p_offensive" in name:
                    cell.number_format = "0.00"
                cells.append(cell)
            ws.row_dimensions[n].height = 64
            ws.append(cells)

    human_col = get_column_letter(MAIN_COLUMNS.index("human_label") + 1)
    validation = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
    validation.errorTitle, validation.error = "Invalid label", "Use 0 or 1; leave blank if unresolved."
    validation.showErrorMessage = True
    validation.add(f"{human_col}2:{human_col}{max(2, len(rows) + 1)}")
    main.data_validations.append(validation)
    put(main, MAIN_COLUMNS, rows, len(rows))
    put(voting, VOTE_COLUMNS, votes, len(votes))
    summary_rows = [{"metric": k, "value": v} for k, v in summary.items()]
    summary_rows += [{"metric": "excel_truncated_cells", "value": truncations},
                     {"metric": "0 / 1 / blank", "value": "0 = non-offensive; 1 = offensive; blank = unresolved, never 0"},
                     {"metric": "orange", "value": "Review required; quotations/education and other ambiguity are blocked by default"},
                     {"metric": "yellow", "value": "Tie or model abstention; final label blank"},
                     {"metric": "red", "value": "Failed or pending model votes; final label blank"},
                     {"metric": "green", "value": "Imported human review; not automatically a double-annotated gold label"},
                     {"metric": "review fields", "value": "Fill human_label, reviewer and optional review_note; run review command on a COPY of this file."},
                     {"metric": "clean.csv", "value": "Includes finalized model-consensus silver labels and human-reviewed labels; inspect quality_tier."},
                     {"metric": "privacy", "value": "This workbook contains original text; keep it private."}]
    put(info, ["metric", "value"], summary_rows, len(summary_rows))
    fd, temporary = tempfile.mkstemp(dir=Path(path).parent, suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(temporary)
        os.replace(temporary, path)
    finally:
        wb.close()
        if os.path.exists(temporary):
            os.unlink(temporary)


def import_reviews(directory, workbook_path):
    from openpyxl import load_workbook
    from .pipeline import collect_results, load_run
    from .storage import RunLock, Store
    with RunLock(directory):
        config, records, manifest = load_run(directory)
        known = {r["id"]: r for r in records}
        if Path(workbook_path).resolve() == (Path(directory) / "results.xlsx").resolve():
            raise ValueError("Review a copy (e.g. reviewed.xlsx); export overwrites results.xlsx")
        wb = load_workbook(workbook_path, read_only=True, data_only=False)
        updates, seen = {}, set()
        try:
            ws = wb["Results"]
            iterator = ws.iter_rows()
            headers = [c.value for c in next(iterator)]
            required = {"id", "content_hash", "run_id", "human_label", "reviewer", "review_note", "text", "context"}
            if len(headers) != len(set(headers)) or not required.issubset(headers):
                raise ValueError("Review workbook columns missing or duplicated")
            for cells in iterator:
                row = dict(zip(headers, [c.value for c in cells]))
                if not any(v is not None for v in row.values()):
                    continue
                if any(c.data_type == "f" for c in cells):
                    raise ValueError("Review workbook contains formula cells")
                ident = str(row["id"])
                if ident not in known or ident in seen:
                    raise ValueError("Unknown or duplicate review ID")
                seen.add(ident)
                record = known[ident]
                if row["run_id"] != manifest["run_id"] or row["content_hash"] != record["content_hash"] or row["text"] != excel_text(record["text"]) or (row["context"] or "") != excel_text(record["context"]):
                    raise ValueError(f"Review provenance or original text changed for {ident}")
                label = row["human_label"]
                if label is None or label == "":
                    continue
                if type(label) is bool or label not in (0, 1, "0", "1"):
                    raise ValueError(f"Invalid human label for {ident}")
                reviewer = row["reviewer"]
                if not isinstance(reviewer, str) or not reviewer.strip():
                    raise ValueError(f"Reviewer required for {ident}")
                updates[ident] = {"human_label": int(label), "reviewer": reviewer.strip(), "review_note": str(row["review_note"] or "")}
        finally:
            wb.close()
        store = Store(directory)
        try:
            store.set_reviews(updates)
            export_run(store, collect_results(store, config, records, manifest), manifest)
        finally:
            store.close()
        return len(updates)
